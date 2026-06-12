"""Research workbook import/export helpers for address evaluation."""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.scoring import build_research_required_fields, build_research_tables


STATUS_TO_LABEL = {
    "included": "计入",
    "pending_review": "待核验",
    "excluded": "误匹配",
    "manual_added": "人工新增",
}
LABEL_TO_STATUS = {
    "计入": "included",
    "待核验": "pending_review",
    "误匹配": "excluded",
    "误匹配排除": "excluded",
    "人工新增": "manual_added",
    "人工补充": "manual_added",
}

BASE_COLUMNS = ["名称", "类型", "距离(m)", "地址/依据", "状态", "来源", "需补充字段", "备注"]
COMPETITOR_COLUMNS = BASE_COLUMNS + ["配置", "机器数", "面积(㎡)", "小时价(元)", "套餐价", "上座率(%)", "开业年限", "月售", "年售", "充值信息", "置信度"]
FACILITY_COLUMNS = BASE_COLUMNS + ["营业时间", "开业年限", "是否营业到凌晨", "是否24小时", "摊位数量", "规模"]
PROPERTY_COLUMNS = ["字段", "值", "说明"]

REQUIRED_SHEETS = [
    "竞品",
    "餐饮",
    "夜市摊",
    "娱乐配套",
    "便利店",
    "停车场",
    "教育客群",
    "政策红线",
    "物业条件",
    "商圈容量参数",
    "缺失字段说明",
]

PROPERTY_FIELD_MAP = {
    "月租金（元/月）": "monthly_rent",
    "面积（㎡）": "area_sqm",
    "楼层（层）": "floor",
    "门头可见性": "frontage_visibility",
    "停车便利": "parking_convenience",
    "消防满足": "fire_safety_ready",
    "电力容量满足": "power_capacity_ready",
    "空调/排烟满足": "hvac_ready",
    "物业限制": "property_restriction",
}

CAPACITY_FIELD_MAP = {
    "18-35岁有效人口": "effective_population_18_35",
    "流动人口（人/月）": "floating_population",
    "转化率（%）": "conversion_rate_pct",
    "月均消费频次": "monthly_frequency",
    "客单价（元）": "avg_spend",
    "健康月营收（元/月）": "healthy_monthly_revenue",
}


def build_research_workbook(record) -> BytesIO:
    manual_data = record.manual_data or {}
    research_status = build_research_required_fields(manual_data)
    research_tables = build_research_tables(record.dimensions or {}, manual_data)

    wb = Workbook()
    wb.remove(wb.active)
    _write_instruction_sheet(wb, record, research_status)
    _write_table_sheet(wb, "竞品", COMPETITOR_COLUMNS, _table_rows(research_tables, "competitors"), "配置、机器数、面积、小时价、上座率、开业年限、月售/年售、充值信息")
    _write_table_sheet(wb, "餐饮", FACILITY_COLUMNS, _table_rows(research_tables, "food_places"), "营业时间、是否营业到凌晨、开业年限")
    _write_table_sheet(wb, "夜市摊", FACILITY_COLUMNS, _table_rows(research_tables, "night_markets"), "摊位数量、距离、营业时间、规模")
    _write_table_sheet(wb, "娱乐配套", FACILITY_COLUMNS, _table_rows(research_tables, "entertainment_places"), "营业时间、开业年限、类型")
    _write_table_sheet(wb, "便利店", FACILITY_COLUMNS, _table_rows(research_tables, "convenience_stores"), "是否24小时、营业时间")
    _write_table_sheet(wb, "停车场", BASE_COLUMNS, _table_rows(research_tables, "parking_places"), "停车便利性、车位规模、收费情况")
    _write_table_sheet(wb, "教育客群", BASE_COLUMNS, _table_rows(research_tables, "education"), "核验是否为有效大学/中职/技校客群")
    _write_table_sheet(wb, "政策红线", BASE_COLUMNS, _table_rows(research_tables, "policy_redline"), "核验小学、幼儿园、中学、政府机构 200m 红线")
    _write_kv_sheet(wb, "物业条件", PROPERTY_FIELD_MAP, manual_data.get("property_conditions") or manual_data)
    _write_kv_sheet(wb, "商圈容量参数", CAPACITY_FIELD_MAP, manual_data.get("market_capacity_inputs") or manual_data)
    _write_missing_sheet(wb, research_status)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def parse_research_workbook(file_obj) -> dict[str, Any]:
    wb = load_workbook(file_obj, data_only=True)
    errors: list[dict[str, Any]] = []
    missing_sheets = [name for name in REQUIRED_SHEETS if name not in wb.sheetnames]
    for sheet_name in missing_sheets:
        errors.append({"sheet": sheet_name, "row": None, "message": f"缺少 Sheet：{sheet_name}"})

    manual_data: dict[str, Any] = {
        "competitors": [],
        "food_places": [],
        "night_markets": [],
        "entertainment_places": [],
        "convenience_stores": [],
        "property_conditions": {},
        "market_capacity_inputs": {},
    }

    sheet_targets = {
        "竞品": ("competitors", COMPETITOR_COLUMNS),
        "餐饮": ("food_places", FACILITY_COLUMNS),
        "夜市摊": ("night_markets", FACILITY_COLUMNS),
        "娱乐配套": ("entertainment_places", FACILITY_COLUMNS),
        "便利店": ("convenience_stores", FACILITY_COLUMNS),
    }
    for sheet_name, (target, _columns) in sheet_targets.items():
        if sheet_name in wb.sheetnames:
            manual_data[target], sheet_errors = _parse_rows_sheet(wb[sheet_name], sheet_name)
            errors.extend(sheet_errors)

    # Parking, education and redline are kept for audit but do not currently affect scoring.
    for sheet_name, target in [("停车场", "parking_places"), ("教育客群", "education_places"), ("政策红线", "policy_redline_review")]:
        if sheet_name in wb.sheetnames:
            manual_data[target], sheet_errors = _parse_rows_sheet(wb[sheet_name], sheet_name)
            errors.extend(sheet_errors)

    if "物业条件" in wb.sheetnames:
        manual_data["property_conditions"], sheet_errors = _parse_kv_sheet(wb["物业条件"], PROPERTY_FIELD_MAP, "物业条件")
        errors.extend(sheet_errors)
    if "商圈容量参数" in wb.sheetnames:
        manual_data["market_capacity_inputs"], sheet_errors = _parse_kv_sheet(wb["商圈容量参数"], CAPACITY_FIELD_MAP, "商圈容量参数")
        errors.extend(sheet_errors)

    summary = {
        "competitors": len(manual_data["competitors"]),
        "food_places": len(manual_data["food_places"]),
        "night_markets": len(manual_data["night_markets"]),
        "entertainment_places": len(manual_data["entertainment_places"]),
        "convenience_stores": len(manual_data["convenience_stores"]),
        "parking_places": len(manual_data.get("parking_places") or []),
        "education_places": len(manual_data.get("education_places") or []),
        "policy_redline_review": len(manual_data.get("policy_redline_review") or []),
        "property_fields": len([v for v in manual_data["property_conditions"].values() if v not in (None, "")]),
        "capacity_fields": len([v for v in manual_data["market_capacity_inputs"].values() if v not in (None, "")]),
    }
    return {
        "success": not errors,
        "manual_data": manual_data,
        "summary": summary,
        "errors": errors,
    }


def _write_instruction_sheet(wb: Workbook, record, research_status: dict) -> None:
    ws = wb.create_sheet("使用说明")
    rows = [
        ["评估地址", record.address],
        ["评估记录ID", record.id],
        ["调研完整度", f"{research_status.get('completion_rate', 0)}%"],
        ["使用流程", "核验高德底表 -> 补齐字段 -> 将状态改为 计入/人工新增 -> 上传系统 -> 重新生成报告"],
        ["状态说明", "计入/人工新增参与评分；待核验/误匹配不参与评分。"],
    ]
    for row in rows:
        ws.append(row)
    _style_sheet(ws)


def _write_table_sheet(wb: Workbook, title: str, columns: list[str], rows: list[dict], default_missing: str) -> None:
    ws = wb.create_sheet(title)
    ws.append(columns)
    for row in rows:
        values = []
        for col in columns:
            values.append(_row_value(row, col, default_missing))
        ws.append(values)
    if not rows:
        ws.append([""] * len(columns))
    _style_sheet(ws)


def _write_kv_sheet(wb: Workbook, title: str, fields: dict[str, str], values: dict) -> None:
    ws = wb.create_sheet(title)
    ws.append(PROPERTY_COLUMNS)
    for label, key in fields.items():
        ws.append([label, values.get(key, ""), "请按真实调研填写；缺失则留空"])
    _style_sheet(ws)


def _write_missing_sheet(wb: Workbook, research_status: dict) -> None:
    ws = wb.create_sheet("缺失字段说明")
    ws.append(["字段", "状态", "说明"])
    for item in research_status.get("items") or []:
        ws.append([item.get("label"), item.get("status"), "已补充可留空；缺失项建议优先调研"])
    _style_sheet(ws)


def _table_rows(research_tables: dict, key: str) -> list[dict]:
    confirmed = (research_tables.get("confirmed") or {}).get(key) or {}
    excluded = research_tables.get("excluded") or {}
    rows = []
    if isinstance(confirmed, dict):
        for part in ("amap", "pending", "manual"):
            value = confirmed.get(part)
            if isinstance(value, list):
                rows.extend(value)
    excluded_rows = excluded.get(key)
    if isinstance(excluded_rows, list):
        rows.extend(excluded_rows)
    return rows


def _row_value(row: dict, col: str, default_missing: str):
    mapping = {
        "名称": row.get("name"),
        "类型": row.get("classification_label") or row.get("type"),
        "距离(m)": row.get("distance"),
        "地址/依据": row.get("classification_reason") or row.get("address") or row.get("notes"),
        "状态": STATUS_TO_LABEL.get(row.get("status"), row.get("status") or "待核验"),
        "来源": _display_source(row.get("source") or row.get("data_source")),
        "需补充字段": row.get("missing_fields") or default_missing,
        "备注": row.get("notes"),
        "配置": row.get("configuration"),
        "机器数": row.get("machine_count"),
        "面积(㎡)": row.get("area_sqm"),
        "小时价(元)": row.get("hourly_price"),
        "套餐价": row.get("package_price"),
        "上座率(%)": row.get("occupancy_rate"),
        "开业年限": row.get("open_years"),
        "月售": row.get("monthly_sales"),
        "年售": row.get("annual_sales"),
        "充值信息": row.get("recharge_info"),
        "置信度": row.get("confidence"),
        "营业时间": row.get("business_hours"),
        "是否营业到凌晨": _bool_label(row.get("late_night")),
        "是否24小时": _bool_label(row.get("is_24h")),
        "摊位数量": row.get("stall_count"),
        "规模": row.get("scale"),
    }
    value = mapping.get(col)
    return "" if value is None else value


def _parse_rows_sheet(ws, sheet_name: str) -> tuple[list[dict], list[dict]]:
    errors = []
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
    header_map = {name: idx for idx, name in enumerate(headers) if name}
    for required in ["名称", "状态"]:
        if required not in header_map:
            errors.append({"sheet": sheet_name, "row": 1, "message": f"缺少列：{required}"})
    rows = []
    for row_idx, row_cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        row = {header: row_cells[idx] if idx < len(row_cells) else None for header, idx in header_map.items()}
        if not any(value not in (None, "") for value in row.values()):
            continue
        name = _text(row.get("名称"))
        if not name:
            continue
        status_label = _text(row.get("状态")) or "待核验"
        status = LABEL_TO_STATUS.get(status_label)
        if not status:
            errors.append({"sheet": sheet_name, "row": row_idx, "message": f"状态值无效：{status_label}，只能填写 计入/待核验/误匹配/人工新增"})
            status = "pending_review"
        item = {
            "name": name,
            "type": _text(row.get("类型")),
            "distance": _to_number(row.get("距离(m)")),
            "address": _text(row.get("地址/依据")),
            "status": status,
            "include": status not in {"pending_review", "excluded"},
            "source": _text(row.get("来源")) or "人工调研",
            "notes": _text(row.get("备注")),
            "configuration": _text(row.get("配置")),
            "machine_count": _to_number(row.get("机器数")),
            "area_sqm": _to_number(row.get("面积(㎡)")),
            "hourly_price": _to_number(row.get("小时价(元)")),
            "package_price": _text(row.get("套餐价")),
            "occupancy_rate": _to_number(row.get("上座率(%)")),
            "open_years": _to_number(row.get("开业年限")),
            "monthly_sales": _to_number(row.get("月售")),
            "annual_sales": _to_number(row.get("年售")),
            "recharge_info": _text(row.get("充值信息")),
            "confidence": _to_number(row.get("置信度")) or 0.8,
            "business_hours": _text(row.get("营业时间")),
            "late_night": _to_bool(row.get("是否营业到凌晨")),
            "is_24h": _to_bool(row.get("是否24小时")),
            "stall_count": _to_number(row.get("摊位数量")),
            "scale": _text(row.get("规模")),
        }
        rows.append({k: v for k, v in item.items() if v not in (None, "")})
    return rows, errors


def _parse_kv_sheet(ws, fields: dict[str, str], sheet_name: str) -> tuple[dict, list[dict]]:
    errors = []
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
    if "字段" not in headers or "值" not in headers:
        return {}, [{"sheet": sheet_name, "row": 1, "message": "缺少列：字段/值"}]
    field_idx = headers.index("字段")
    value_idx = headers.index("值")
    result = {}
    for row_idx, row_cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        label = _text(row_cells[field_idx] if field_idx < len(row_cells) else None)
        if not label:
            continue
        key = fields.get(label)
        if not key:
            errors.append({"sheet": sheet_name, "row": row_idx, "message": f"未知字段：{label}"})
            continue
        result[key] = row_cells[value_idx] if value_idx < len(row_cells) else None
    return result, errors


def _style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="EAF2FF")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="1F3A5F")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col_idx, column in enumerate(ws.columns, start=1):
        max_len = 10
        for cell in column:
            max_len = max(max_len, len(str(cell.value or "")) + 2)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len, 34)


def _display_source(value: Any) -> str:
    if value in {"amap", "api"}:
        return "高德API"
    return str(value or "")


def _bool_label(value: Any) -> str:
    if value is True:
        return "是"
    if value is False:
        return "否"
    return ""


def _to_bool(value: Any):
    text = _text(value).lower()
    if text in {"是", "yes", "true", "1", "y"}:
        return True
    if text in {"否", "no", "false", "0", "n"}:
        return False
    return None


def _to_number(value: Any):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return value
    text = str(value).replace("%", "").replace("m", "").replace("米", "").strip()
    try:
        return float(text)
    except ValueError:
        return None


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()
