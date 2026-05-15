"""
经验文档导入服务。

支持 TXT / DOCX / PDF 文档解析、分块、向量入库和待确认权重建议生成。
"""
import hashlib
import io
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Optional

from sqlalchemy.orm import Session

from app.models.store import KnowledgeDocument, DocumentInsight, ScoringRule, Store
from app.services.importer import UPLOAD_ROOT
from app.services.vector_rag import ensure_vector_table, store_text_as_vector

ALLOWED_DOCUMENT_EXTENSIONS = {".txt", ".docx", ".pdf"}
MAX_DOCUMENT_SIZE = 20 * 1024 * 1024

DIMENSION_META = {
    "traffic": ("交通与人流", "foot_traffic"),
    "competition": ("竞品分析", "competitor_count"),
    "population": ("目标客群", "young_density"),
    "rent": ("租金与成本", "rent_ratio"),
    "facility": ("配套设施", "commercial_density"),
    "policy": ("政策环境", "policy_risk"),
}

DIMENSION_KEYWORDS = {
    "traffic": ["人流", "客流", "地铁", "公交", "交通", "到达", "可达", "主路", "停车", "消费热度", "热力"],
    "competition": ["竞品", "网吧", "电竞馆", "竞争", "饱和", "同业", "分流"],
    "population": ["学生", "大学", "高校", "年轻人", "白领", "社区", "住宅", "办公", "客群", "会员"],
    "rent": ["租金", "房租", "成本", "物业", "押金", "坪效", "面积"],
    "facility": ["餐饮", "便利店", "商场", "停车场", "配套", "夜宵", "外卖", "商业"],
    "policy": ["消防", "证照", "政策", "合规", "限制", "审批", "噪音", "未成年"],
}

NEGATIVE_KEYWORDS = ["不适合", "不建议", "风险", "过高", "偏高", "失败", "不足", "较差", "缺少", "远", "偏", "饱和"]


def validate_document_upload(filename: str, file_size: int) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext not in ALLOWED_DOCUMENT_EXTENSIONS:
        raise ValueError("仅支持 .txt / .docx / .pdf 经验文档")
    if file_size <= 0:
        raise ValueError("文件为空，无法解析")
    if file_size > MAX_DOCUMENT_SIZE:
        raise ValueError("文件大小超过 20MB 限制")
    return ext.lstrip(".")


async def process_document_upload(
    file_bytes: bytes,
    filename: str,
    scope_type: str,
    tenant_id: int,
    user_id: int,
    db: Session,
    store_id: Optional[int] = None,
    candidate_address: Optional[str] = None,
) -> KnowledgeDocument:
    file_type = validate_document_upload(filename, len(file_bytes))
    _validate_scope(scope_type, tenant_id, db, store_id, candidate_address)

    stored_path = _save_document_file(file_bytes, filename, tenant_id)
    document = KnowledgeDocument(
        tenant_id=tenant_id,
        store_id=store_id if scope_type == "store" else None,
        original_filename=os.path.basename(filename or "document"),
        stored_path=stored_path,
        file_type=file_type,
        file_size=len(file_bytes),
        file_hash=hashlib.sha256(file_bytes).hexdigest(),
        scope_type=scope_type,
        candidate_address=candidate_address.strip() if scope_type == "candidate" and candidate_address else None,
        parse_status="processing",
        vector_status="pending",
        uploaded_by=user_id,
    )
    db.add(document)
    db.commit()
    db.refresh(document)

    try:
        text = extract_document_text(file_bytes, file_type)
        if len(text.strip()) < 20:
            raise ValueError("文档可解析文本过少。扫描版 PDF 需要先 OCR 后再上传。")

        chunks = chunk_text(text)
        document.summary = summarize_text(text)
        document.chunks = [{"index": i, "content": chunk[:1200]} for i, chunk in enumerate(chunks)]
        document.chunk_count = len(chunks)
        document.parse_status = "success"
        document.parse_message = f"解析完成，共 {len(chunks)} 个文本块"
        db.add(document)
        db.commit()
        db.refresh(document)

        await _vectorize_document(document, chunks, tenant_id, db)
        _create_document_insights(document, text, tenant_id, db)
        db.commit()
        db.refresh(document)
    except Exception as e:
        db.rollback()
        document = db.query(KnowledgeDocument).filter(KnowledgeDocument.id == document.id).first()
        if document:
            document.parse_status = "failed"
            document.parse_message = str(e)
            document.vector_status = "failed"
            document.vector_message = "文档解析失败，未进入知识库"
            db.add(document)
            db.commit()
            db.refresh(document)

    return document


def extract_document_text(file_bytes: bytes, file_type: str) -> str:
    if file_type == "txt":
        return _decode_text(file_bytes)
    if file_type == "docx":
        try:
            from docx import Document
        except ImportError as e:
            raise RuntimeError("缺少 python-docx 依赖，请先安装 backend/requirements.txt") from e
        doc = Document(io.BytesIO(file_bytes))
        parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if file_type == "pdf":
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise RuntimeError("缺少 pypdf 依赖，请先安装 backend/requirements.txt") from e
        reader = PdfReader(io.BytesIO(file_bytes))
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
        return "\n".join([p for p in pages if p])
    raise ValueError(f"不支持的文档类型: {file_type}")


def chunk_text(text: str, max_chars: int = 1400, overlap: int = 160) -> list[str]:
    paragraphs = [p.strip() for p in re.split(r"\n{2,}|\r\n{2,}", text) if p.strip()]
    chunks: list[str] = []
    current = ""
    for para in paragraphs:
        if len(current) + len(para) + 1 <= max_chars:
            current = f"{current}\n{para}".strip()
        else:
            if current:
                chunks.append(current)
            if len(para) <= max_chars:
                current = para
            else:
                start = 0
                while start < len(para):
                    chunks.append(para[start:start + max_chars])
                    start += max_chars - overlap
                current = ""
    if current:
        chunks.append(current)
    return chunks or [text[:max_chars]]


def summarize_text(text: str, max_chars: int = 500) -> str:
    cleaned = re.sub(r"\s+", " ", text).strip()
    return cleaned[:max_chars] + ("..." if len(cleaned) > max_chars else "")


async def _vectorize_document(document: KnowledgeDocument, chunks: list[str], tenant_id: int, db: Session) -> None:
    await ensure_vector_table(db)
    stored = 0
    for idx, chunk in enumerate(chunks):
        source_id = _stable_source_id(f"document:{document.id}:{idx}")
        vec_id = await store_text_as_vector(
            content=chunk,
            metadata={
                "type": "document_experience",
                "document_id": document.id,
                "filename": document.original_filename,
                "scope_type": document.scope_type,
                "store_id": document.store_id,
                "candidate_address": document.candidate_address,
                "chunk_index": idx,
            },
            source_type="document_experience",
            source_id=source_id,
            tenant_id=tenant_id,
            db=db,
        )
        if vec_id:
            stored += 1

    document.vector_status = "success" if stored == len(chunks) else ("partial" if stored else "failed")
    document.vector_message = f"已入库 {stored}/{len(chunks)} 个文本块"
    db.add(document)
    db.commit()


def _create_document_insights(document: KnowledgeDocument, text: str, tenant_id: int, db: Session) -> None:
    sentences = _split_sentences(text)
    seen: set[tuple[str, str]] = set()
    created = 0
    for sentence in sentences:
        for dimension, keywords in DIMENSION_KEYWORDS.items():
            if not any(kw in sentence for kw in keywords):
                continue
            direction = "decrease" if any(kw in sentence for kw in NEGATIVE_KEYWORDS) else "increase"
            dimension_name, sub_factor = DIMENSION_META[dimension]
            key = (dimension, sentence[:80])
            if key in seen:
                continue
            seen.add(key)
            insight = DocumentInsight(
                tenant_id=tenant_id,
                document_id=document.id,
                dimension=dimension,
                dimension_name=dimension_name,
                sub_factor=sub_factor,
                insight=_build_insight_text(dimension_name, direction, sentence),
                evidence=sentence[:500],
                adjustment_direction=direction,
                suggested_weight=_suggest_weight(db, tenant_id, dimension, sub_factor, direction),
                confidence=_confidence_for_sentence(sentence),
                status="pending",
            )
            db.add(insight)
            created += 1
            if created >= 8:
                return


def approve_document_insight(insight: DocumentInsight, user_id: int, db: Session) -> ScoringRule:
    rule = db.query(ScoringRule).filter(
        ScoringRule.tenant_id == insight.tenant_id,
        ScoringRule.dimension == insight.dimension,
        ScoringRule.sub_factor == insight.sub_factor,
        ScoringRule.is_active == True,
    ).first()
    if not rule:
        raise ValueError("未找到可映射的评分规则，无法生效")

    current = rule.effective_weight or rule.dynamic_weight or rule.base_weight
    target = insight.suggested_weight if insight.suggested_weight is not None else current
    bounded = max(0.01, min(0.40, target))
    if abs(bounded - current) > 0.03:
        bounded = current + (0.03 if bounded > current else -0.03)

    rule.dynamic_weight = round(bounded, 4)
    rule.effective_weight = round(bounded, 4)
    rule.last_updated_by = "document_review"
    rule.update_reason = f"经验文档建议已确认：{insight.insight}；依据：{insight.evidence or ''}"[:1000]
    rule.update_count = (rule.update_count or 0) + 1
    rule.updated_at = datetime.utcnow()

    insight.status = "approved"
    insight.reviewed_by = user_id
    insight.reviewed_at = datetime.utcnow()
    db.add(rule)
    db.add(insight)
    db.commit()
    db.refresh(rule)
    return rule


def reject_document_insight(insight: DocumentInsight, user_id: int, db: Session, note: str = "") -> None:
    insight.status = "rejected"
    insight.reviewed_by = user_id
    insight.reviewed_at = datetime.utcnow()
    insight.review_note = note[:500] if note else None
    db.add(insight)
    db.commit()


def _save_document_file(file_bytes: bytes, filename: str, tenant_id: int) -> str:
    upload_dir = Path(UPLOAD_ROOT) / str(tenant_id) / "documents"
    upload_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    safe_name = os.path.basename(filename or "document")
    path = upload_dir / f"{timestamp}_{safe_name}"
    path.write_bytes(file_bytes)
    return str(path)


def _validate_scope(scope_type: str, tenant_id: int, db: Session, store_id: Optional[int], candidate_address: Optional[str]) -> None:
    if scope_type not in {"brand", "store", "candidate"}:
        raise ValueError("scope_type 仅支持 brand / store / candidate")
    if scope_type == "store":
        if not store_id:
            raise ValueError("绑定门店文档必须选择门店")
        store = db.query(Store).filter(Store.id == store_id, Store.tenant_id == tenant_id).first()
        if not store:
            raise ValueError("绑定门店不存在")
    if scope_type == "candidate" and not (candidate_address or "").strip():
        raise ValueError("候选地址文档必须填写候选地址")


def _decode_text(file_bytes: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore")


def _split_sentences(text: str) -> list[str]:
    cleaned = re.sub(r"\s+", " ", text)
    parts = re.split(r"[。！？!?；;]\s*", cleaned)
    return [p.strip() for p in parts if len(p.strip()) >= 12]


def _build_insight_text(dimension_name: str, direction: str, sentence: str) -> str:
    action = "提高关注权重" if direction == "increase" else "降低或谨慎评估该因素"
    return f"{dimension_name}：建议{action}。文档依据：{sentence[:120]}"


def _confidence_for_sentence(sentence: str) -> float:
    score = 0.55
    if any(word in sentence for word in ["数据", "复盘", "调研", "统计", "验证"]):
        score += 0.15
    if any(word in sentence for word in ["成功", "失败", "营收", "客流", "转化"]):
        score += 0.15
    return min(0.9, score)


def _suggest_weight(db: Session, tenant_id: int, dimension: str, sub_factor: str, direction: str) -> Optional[float]:
    rule = db.query(ScoringRule).filter(
        ScoringRule.tenant_id == tenant_id,
        ScoringRule.dimension == dimension,
        ScoringRule.sub_factor == sub_factor,
        ScoringRule.is_active == True,
    ).first()
    if not rule:
        return None
    current = rule.effective_weight or rule.dynamic_weight or rule.base_weight
    delta = 0.03 if direction == "increase" else -0.03
    return round(max(0.01, min(0.40, current + delta)), 4)


def _stable_source_id(value: str) -> int:
    return int(hashlib.md5(value.encode("utf-8")).hexdigest()[:8], 16) % (2**31 - 1)
